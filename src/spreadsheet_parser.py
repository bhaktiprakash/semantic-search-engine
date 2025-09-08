import openpyxl
import networkx as nx
import re


CELL_REF_REGEX = re.compile(r"([A-Z]+[0-9]+)")

def build_dependency_graph(file_path: str):
    """
    Parses an Excel file and builds a directed graph of formula dependencies.
    An edge from A to B means B's formula depends on A's value.
    """
    try:
        workbook = openpyxl.load_workbook(file_path, data_only=False) 
    except FileNotFoundError:
        print(f"Error: File not found at {file_path}")
        return None, None

    graph = nx.DiGraph()
    
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                node_id = f"{sheet_name}!{cell.coordinate}"
                graph.add_node(node_id, sheet=sheet_name, coordinate=cell.coordinate, value=cell.value, formula=None)
                
                if cell.data_type == 'f': 
                    formula = cell.value
                    graph.nodes[node_id]['formula'] = formula
                    
                    
                    dependencies = CELL_REF_REGEX.findall(formula)
                    for dep in dependencies:
                        dep_node_id = f"{sheet_name}!{dep}"
                        graph.add_edge(dep_node_id, node_id)

    print(f"Built dependency graph for {file_path} with {graph.number_of_nodes()} nodes and {graph.number_of_edges()} edges.")
    return workbook, graph

